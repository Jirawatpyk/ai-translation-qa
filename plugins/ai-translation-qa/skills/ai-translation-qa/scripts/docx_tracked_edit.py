#!/usr/bin/env python3
"""Apply tracked changes (Word revisions) to a .docx — a reusable helper for
vendor QA / source-check deliverables where corrections must be visible and
acceptable/rejectable in Word.

Two ways to use it:

1) As a library (recommended — you control what to change):

    from docx import Document
    from docx_tracked_edit import (
        cell_text, tracked_replace_paragraph, set_cell_text,
        enable_track_changes, verify_roundtrip)

    doc = Document("in.docx")
    # replace text in a table cell as a tracked del+ins, preserving run format:
    cell = doc.tables[7].rows[3].cells[1]
    tracked_replace_paragraph(cell.paragraphs[0],
                              old="wrong", new="right", author="QA", date="2026-08-04T00:00:00Z")
    enable_track_changes(doc)                 # open in Word with revisions ON
    doc.save("out.docx")
    verify_roundtrip("out.docx", ...)         # see below

2) As a CLI, applying a JSON edit list:

    python docx_tracked_edit.py in.docx edits.json out.docx --author "QA"

    edits.json: [{"table":7,"row":3,"col":1,"old":"wrong","new":"right"}, ...]
    (col defaults to 1 — the Source column in an Annotation|Source|Target table)
    Exit codes: 0 = all edits applied · 2 = some edits skipped (by-design
    refusal; inspect stderr — the output file WAS written) · 1 = uncaught
    invocation failure. Caveat shared with the XLIFF I/O scripts: argparse
    itself also exits 2 on a malformed command line; distinguish by whether
    the output file exists.

GOTCHA THIS SCRIPT SOLVES: python-docx's `cell.text` / `paragraph.text` do NOT see
runs nested inside <w:ins> (tracked insertions). After a tracked del+ins replacement,
the original runs live in <w:del> and the new text in <w:ins>, so `.text` reads back
'' (empty) and your edit looks like it vanished. Use `cell_text()` / `para_text()`
here, which walk the XML and honor ins/del, for any read-back or verification.

SCOPE / LIMITATIONS (this tool is built for a single, clean pass over a freshly
extracted file — the standard vendor-QA / source-check workflow):
- It refuses to edit a paragraph it cannot rewrite safely, rather than corrupting it:
  one that ALREADY contains tracked changes (text would duplicate), one whose text
  lives in a hyperlink or field (same duplication — python-docx's p.runs cannot see
  that text), or one holding an image (accepting the revision would delete it). Via
  the CLI/apply_edits these become per-edit "skipped" entries; the library call
  raises. Accept/reject existing revisions first, or edit a clean copy.
- A tracked insertion carries the FIRST run's formatting; replacing a span of mixed
  formatting collapses it to that one format in accept-view. Fine for plain text.
- `apply_edits` composes edits per cell with successive string replacement, which is
  global (replaces every occurrence) and order-sensitive. Give one edit per distinct
  `old` string; don't chain edits where a later `old` could match earlier `new` text.
"""
import argparse, copy, json, sys
from datetime import datetime, timezone
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_counter = [1000]
def _nid():
    _counter[0] += 1
    return str(_counter[0])

def _today():
    """Current UTC date stamp for revisions — a real audit trail, not a fixed date."""
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT00:00:00Z")

def _para_text(p_el):
    """Accept-view text of ONE <w:p>, honoring ins/del and <w:br>/<w:tab>.
    (Accept-view only: tracked insertions are included, deletions skipped.
    A reject-view reader would be a separate function — do not overload this
    one with view flags nobody passes.)"""
    out = []
    for node in p_el.iter():
        anc = {a.tag for a in node.iterancestors()}
        in_del = qn('w:del') in anc
        if node.tag == qn('w:t'):
            if in_del:  # deleted text: skip in accept-view
                continue
            out.append(node.text or "")
        elif node.tag == qn('w:br'):
            if not in_del:
                out.append("\n")
        elif node.tag == qn('w:tab'):
            if not in_del:
                out.append("\t")
    return "".join(out)

def cell_text(cell):
    """Accept-view text of a table cell, INCLUDING tracked insertions. Paragraph
    boundaries within the cell become '\\n' (matching python-docx's cell.text)."""
    paras = cell._tc.findall(qn('w:p'))
    return "\n".join(_para_text(p) for p in paras)

def para_text(p):
    """Accept-view text of a paragraph, INCLUDING tracked insertions."""
    return _para_text(p._p)

def _has_tracked(p_el):
    return p_el.find(qn('w:ins')) is not None or p_el.find(qn('w:del')) is not None

def _edit_blocker(p):
    """Why this paragraph must not be tracked-edited, or None if it is safe.

    Every case here would corrupt the file rather than fail loudly, so the check
    is mandatory before any replacement — do not regress by skipping it:
      * already-tracked: p.runs sees only direct-child runs, so content already
        inside w:ins/w:del would not be wrapped and the text duplicates.
      * hyperlink / field: the visible text lives in w:hyperlink or w:fldSimple,
        which p.runs cannot reach either — the deletion covers only part of the
        text and the insertion adds the whole of it back, duplicating the rest.
      * image: a w:drawing lives inside a run, so wrapping the runs in w:del
        deletes the picture when the revision is accepted."""
    p_el = p._p
    if _has_tracked(p_el):
        return ("paragraph already contains tracked changes; accept/reject them "
                "before editing again (this tool is for a single clean pass)")
    if p_el.find('.//' + qn('w:drawing')) is not None:
        return ("paragraph contains an image (w:drawing) — a tracked replacement "
                "would delete the image when the revision is accepted")
    if para_text(p) != "".join(r.text for r in p.runs):
        return "paragraph contains hyperlink/field content"
    return None

def _append_line(run_el, line):
    """Append one line of text to a <w:r>, turning tabs into real <w:tab/>.
    A literal tab inside <w:t> is collapsed to a space by Word."""
    wrote = False
    for i, chunk in enumerate(line.split("\t")):
        if i:
            run_el.append(OxmlElement('w:tab')); wrote = True
        if chunk:
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = chunk
            run_el.append(t); wrote = True
    return wrote

def tracked_replace_paragraph(p, old=None, new=None, author="QA", date=None):
    """Replace the paragraph's text with a tracked deletion + insertion.
    If `old` is given, only replaces when `old` is in the current text (and the new
    text is current-with-old-swapped); otherwise replaces the whole paragraph text
    with `new`. Returns True if a change was written.

    Raises ValueError when the paragraph cannot be edited safely — it already
    contains tracked changes, or its text lives in a hyperlink/field, or it holds an
    image. See _edit_blocker: each of those would corrupt the paragraph silently, so
    the refusal is deliberate. apply_edits() turns the same conditions into a per-edit
    skip instead of raising."""
    if date is None:
        date = _today()
    blocked = _edit_blocker(p)
    if blocked:
        raise ValueError(blocked)
    cur = para_text(p)
    if old is not None:
        if old not in cur:
            return False
        target = cur.replace(old, new)
    else:
        target = new
    if target == cur:
        return False
    runs = p.runs
    rPr = runs[0]._r.find(qn('w:rPr')) if runs else None
    # wrap all existing runs as a tracked deletion
    dele = OxmlElement('w:del')
    dele.set(qn('w:id'), _nid()); dele.set(qn('w:author'), author); dele.set(qn('w:date'), date)
    for r in list(runs):
        r_el = r._r
        for t in r_el.findall(qn('w:t')):
            dt = OxmlElement('w:delText'); dt.set(qn('xml:space'), 'preserve')
            dt.text = t.text; r_el.replace(t, dt)
        r_el.getparent().remove(r_el); dele.append(r_el)
    # tracked insertion of the new text (\n -> <w:br>, \t -> <w:tab/>)
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), _nid()); ins.set(qn('w:author'), author); ins.set(qn('w:date'), date)
    nr = OxmlElement('w:r')
    if rPr is not None:
        nr.append(copy.deepcopy(rPr))
    wrote = False
    for i, line in enumerate(target.split("\n")):
        if i:
            nr.append(OxmlElement('w:br')); wrote = True
        wrote = _append_line(nr, line) or wrote
    if not wrote:
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = ""
        nr.append(t)
    ins.append(nr)
    pPr = p._p.find(qn('w:pPr'))
    idx = list(p._p).index(pPr) + 1 if pPr is not None else 0
    p._p.insert(idx, ins); p._p.insert(idx, dele)
    return True

def set_cell_text(cell, text):
    """Overwrite a cell's text WITHOUT tracking (use for filling an empty Target
    cell), keeping the cell's first-run formatting. \\n becomes a line break.
    Also strips any pre-existing tracked ins/del in the cell so leftover revisions
    can't survive alongside the new text.

    Raises ValueError when the cell holds content this function cannot clear:
    text inside a hyperlink or field, or an image. Clearing only the direct
    runs would leave that content in place NEXT TO the new text — old and new
    side by side in the deliverable, with no error. Same refuse-don't-corrupt
    contract as tracked_replace_paragraph."""
    for bad_tag, why in ((qn('w:hyperlink'), 'hyperlink text'),
                         (qn('w:fldSimple'), 'field text'),
                         (qn('w:instrText'), 'field text'),
                         (qn('w:drawing'), 'an image')):
        if cell._tc.find('.//' + bad_tag) is not None:
            raise ValueError(
                'cell contains %s, which set_cell_text cannot clear — the old '
                'content would survive next to the new text; edit this cell '
                'manually' % why)
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    rPr = None
    # capture formatting from a direct run or one nested in an ins, then clear both
    for r in p._p.findall(qn('w:r')):
        if rPr is None:
            rPr = r.find(qn('w:rPr'))
        p._p.remove(r)
    for wrap in p._p.findall(qn('w:ins')) + p._p.findall(qn('w:del')):
        if rPr is None:
            rr = wrap.find(qn('w:r'))
            if rr is not None:
                rPr = rr.find(qn('w:rPr'))
        p._p.remove(wrap)
    for i, line in enumerate(text.split("\n")):
        r = OxmlElement('w:r')
        if rPr is not None:
            r.append(copy.deepcopy(rPr))
        if i:
            r.append(OxmlElement('w:br'))
        if not _append_line(r, line) and not i:
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = ""
            r.append(t)
        p._p.append(r)

def enable_track_changes(doc):
    """Make Word open the file with Track Changes turned ON."""
    st = doc.settings.element
    if st.find(qn('w:trackChanges')) is None:
        st.insert(0, OxmlElement('w:trackChanges'))

def apply_edits(doc, edits, author="QA", date=None):
    """Apply a list of {table,row,col,old,new} edits. Groups by cell so multiple
    swaps in one paragraph compose.

    Returns (applied, failed). `applied` holds the edits that landed; `failed`
    holds a copy of each edit that did not, with a "skipped" key naming the
    reason. A paragraph this tool cannot edit safely (hyperlink/field text, an
    image, pre-existing revisions) is SKIPPED with that reason rather than
    edited — writing it would corrupt the paragraph.

    Edits are tracked by INDEX, not by value: the same edit can match in several
    paragraphs, and counting each match would report more applied edits than were
    submitted."""
    if date is None:
        date = _today()
    from collections import defaultdict
    byrow = defaultdict(list)
    for i, e in enumerate(edits):
        byrow[(e["table"], e["row"], e.get("col", 1))].append((i, e))
    applied_idx, skips = set(), {}
    for (ti, ri, ci), group in byrow.items():
        cell = doc.tables[ti].rows[ri].cells[ci]
        for p in cell.paragraphs:
            cur = para_text(p); new = cur
            used = []
            for i, e in group:
                if e["old"] in new:
                    new = new.replace(e["old"], e["new"]); used.append(i)
            if not used or new == cur:
                continue
            blocked = _edit_blocker(p)
            if blocked:
                for i in used:
                    skips.setdefault(i, blocked)
                continue
            if tracked_replace_paragraph(p, new=new, author=author, date=date):
                applied_idx.update(used)
    applied = [edits[i] for i in sorted(applied_idx)]
    failed = []
    for i, e in enumerate(edits):
        if i in applied_idx:
            continue
        f = dict(e)
        f["skipped"] = skips.get(
            i, "no paragraph in the target cell contains this 'old' text")
        failed.append(f)
    return applied, failed

def verify_roundtrip(path, expected):
    """expected: {(table,row,col): accepted_text}. Returns list of mismatches
    using tracked-aware reads. Empty list = every cell matches."""
    doc = Document(path); bad = []
    for (ti, ri, ci), want in expected.items():
        got = cell_text(doc.tables[ti].rows[ri].cells[ci])
        if got.strip() != want.strip():
            bad.append((ti, ri, ci, want[:40], got[:40]))
    return bad

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("infile"); ap.add_argument("edits"); ap.add_argument("outfile")
    ap.add_argument("--author", default="QA")
    ap.add_argument("--date", default=None, help="revision date; defaults to today (UTC)")
    a = ap.parse_args()
    doc = Document(a.infile)
    edits = json.load(open(a.edits, encoding="utf-8"))
    applied, failed = apply_edits(doc, edits, a.author, a.date or _today())
    enable_track_changes(doc)
    doc.save(a.outfile)
    print(f"applied {len(applied)} / {len(edits)} tracked edits -> {a.outfile}", file=sys.stderr)
    if failed:
        print(f"NOT APPLIED ({len(failed)}): {failed}", file=sys.stderr)
        # Exit 2, matching the XLIFF siblings' convention: 2 = "some edits did
        # not land" (a by-design refusal; the output file was still written).
        # NB: argparse also exits 2 on a malformed command line — a caller
        # distinguishing the cases checks whether the output file exists.
        sys.exit(2)

if __name__ == "__main__":
    main()
